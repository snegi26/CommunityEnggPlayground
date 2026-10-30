"""Regenerate the policy fixtures: one policy as DOCX, PDF and Markdown, plus a scanned PDF.

Run with ``uv run python tests/fixtures/policies/generate.py``. The PDF is written by hand
(no extra dependency) with Helvetica text: headings at 14pt, body at 11pt.
"""

import textwrap
from pathlib import Path

import docx

HERE = Path(__file__).parent

# (heading, [paragraphs]) — three headings, as the U5 test scenarios require.
POLICY: list[tuple[str, list[str]]] = [
    (
        "Purpose",
        [
            "This policy defines how Acme grants, reviews, and removes access to production "
            "systems and company data.",
            "It applies to all employees, contractors, and service accounts.",
        ],
    ),
    (
        "Multi-factor Authentication",
        [
            "MFA is required on all production accounts, including administrator and "
            "break-glass accounts.",
            "Hardware security keys are required for members of the infrastructure team.",
        ],
    ),
    (
        "Access Reviews",
        [
            "Access to production systems is reviewed every quarter by the system owner.",
            "Accounts of departed staff are disabled within one business day of their last "
            "working day.",
        ],
    ),
]


def write_markdown(path: Path) -> None:
    lines: list[str] = []
    for heading, paragraphs in POLICY:
        lines += [f"# {heading}", ""]
        for paragraph in paragraphs:
            lines += [*textwrap.wrap(paragraph, 72), ""]
    path.write_text("\n".join(lines))


def write_docx(path: Path) -> None:
    document = docx.Document()
    for heading, paragraphs in POLICY:
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    document.save(str(path))


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _pdf(content: bytes, *, with_font: bool) -> bytes:
    resources = b"<< /Font << /F1 5 0 R >> >>" if with_font else b"<< >>"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources "
        + resources
        + b" /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


def write_pdf(path: Path) -> None:
    ops: list[str] = []
    y = 740.0
    for heading, paragraphs in POLICY:
        y -= 10
        ops.append(f"BT /F1 14 Tf 72 {y:.0f} Td ({_pdf_escape(heading)}) Tj ET")
        y -= 24
        for paragraph in paragraphs:
            for line in textwrap.wrap(paragraph, 80):
                ops.append(f"BT /F1 11 Tf 72 {y:.0f} Td ({_pdf_escape(line)}) Tj ET")
                y -= 14
            y -= 10
    path.write_bytes(_pdf("\n".join(ops).encode("latin-1"), with_font=True))


def write_scanned_pdf(path: Path) -> None:
    # A page that only paints shapes, like a scanned image: no extractable text.
    path.write_bytes(_pdf(b"0.5 g 72 72 468 648 re f", with_font=False))


def main() -> None:
    write_markdown(HERE / "access_control.md")
    write_docx(HERE / "access_control.docx")
    write_pdf(HERE / "access_control.pdf")
    write_scanned_pdf(HERE / "scanned.pdf")


if __name__ == "__main__":
    main()
